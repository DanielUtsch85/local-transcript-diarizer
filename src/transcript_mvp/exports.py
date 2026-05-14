from __future__ import annotations

import json
from html import escape
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt

from .models import SpeakerMapping, SpeakerSegment, TranscriptSegment


def build_docx(segments: list[TranscriptSegment], mapping: SpeakerMapping, title: str) -> bytes:
    document = Document()
    document.add_heading(title or "Transkript", level=1)

    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10.5)

    for segment in segments:
        paragraph = document.add_paragraph()
        speaker_run = paragraph.add_run(f"{segment.timestamp} - {mapping.label_for(segment.speaker)}")
        speaker_run.bold = True
        speaker_run.add_break(WD_BREAK.LINE)
        paragraph.add_run(segment.text)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_html(segments: list[TranscriptSegment], mapping: SpeakerMapping, title: str) -> str:
    rows = []
    for segment in segments:
        rows.append(
            "<section>"
            f"<h2>{escape(segment.timestamp)} - {escape(mapping.label_for(segment.speaker))}</h2>"
            f"<p>{escape(segment.text)}</p>"
            "</section>"
        )
    body = "\n".join(rows)
    return f"""<!doctype html>
<html lang="de">
<head>
  <meta charset="utf-8">
  <title>{escape(title or "Transkript")}</title>
  <style>
    body {{ font-family: Arial, sans-serif; max-width: 820px; margin: 40px auto; line-height: 1.5; }}
    h1 {{ font-size: 28px; }}
    h2 {{ font-size: 15px; margin: 24px 0 6px; }}
    p {{ margin: 0; }}
    section {{ border-bottom: 1px solid #ddd; padding-bottom: 14px; }}
  </style>
</head>
<body>
  <h1>{escape(title or "Transkript")}</h1>
  {body}
</body>
</html>
"""


def build_diarization_json(
    segments: list[TranscriptSegment] | list[SpeakerSegment],
    source_file: str | Path,
    mapping: SpeakerMapping | None = None,
) -> str:
    rows = _diarization_rows(segments, source_file, mapping)
    return json.dumps(rows, ensure_ascii=False, indent=2) + "\n"


def build_diarization_jsonl(
    segments: list[TranscriptSegment] | list[SpeakerSegment],
    source_file: str | Path,
    mapping: SpeakerMapping | None = None,
) -> str:
    rows = _diarization_rows(segments, source_file, mapping)
    return "\n".join(json.dumps(row, ensure_ascii=False) for row in rows) + ("\n" if rows else "")


def _diarization_rows(
    segments: list[TranscriptSegment] | list[SpeakerSegment],
    source_file: str | Path,
    mapping: SpeakerMapping | None,
) -> list[dict]:
    source = str(source_file)
    return [
        {
            "speaker": mapping.label_for(segment.speaker) if mapping else segment.speaker,
            "start": float(segment.start),
            "end": float(segment.end),
            "source_file": source,
        }
        for segment in segments
        if segment.start is not None and segment.end is not None and segment.end > segment.start
    ]
